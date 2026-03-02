import os
import csv
from typing import Dict

from openpyxl import load_workbook


def _norm(s: str) -> str:
    return (s or "").strip()


def load_aliases_from_csv(csv_path: str) -> Dict[str, str]:
    aliases: Dict[str, str] = {}
    if not os.path.exists(csv_path):
        return aliases

    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        # Stöd både gamla och nya rubriker
        for r in reader:
            alias = _norm(r.get("alias_name") or r.get("alias") or r.get("from") or "")
            target = _norm(r.get("author_name") or r.get("author") or r.get("to") or "")
            if not alias or not target:
                continue
            aliases[alias] = target
    return aliases


def load_aliases_from_xlsx(xlsx_path: str, sheet_name: str = "aliases") -> Dict[str, str]:
    aliases: Dict[str, str] = {}
    if not os.path.exists(xlsx_path):
        return aliases

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

    # Förväntar rubriker i rad 1: alias_name | author_name | comment (comment valfri)
    # Vi letar kolumnerna dynamiskt så du kan flytta dem.
    headers = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if v is None:
            continue
        headers[str(v).strip().lower()] = col

    alias_col = headers.get("alias_name")
    author_col = headers.get("author_name")

    if not alias_col or not author_col:
        raise ValueError(
            f"Excel-filen saknar rubrikerna 'alias_name' och/eller 'author_name' på rad 1. "
            f"Hittade: {list(headers.keys())}"
        )

    for row in range(2, ws.max_row + 1):
        alias = _norm(str(ws.cell(row=row, column=alias_col).value or ""))
        target = _norm(str(ws.cell(row=row, column=author_col).value or ""))
        if not alias or not target:
            continue
        aliases[alias] = target

    return aliases


def export_aliases_to_csv(aliases: Dict[str, str], csv_path: str) -> None:
    os.makedirs(os.path.dirname(csv_path) or ".", exist_ok=True)
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["alias_name", "author_name"])
        for alias, target in sorted(aliases.items(), key=lambda x: x[0].lower()):
            w.writerow([alias, target])


def load_aliases_prefer_xlsx(
    xlsx_path: str = "out/author_aliases.xlsx",
    csv_path: str = "out/author_aliases.csv",
) -> Dict[str, str]:
    """
    1) Om out/author_aliases.xlsx finns: läs den och skriv om out/author_aliases.csv automatiskt.
    2) Annars: läs out/author_aliases.csv (som tidigare).
    """
    if os.path.exists(xlsx_path):
        aliases = load_aliases_from_xlsx(xlsx_path)
        export_aliases_to_csv(aliases, csv_path)
        return aliases

    return load_aliases_from_csv(csv_path)
